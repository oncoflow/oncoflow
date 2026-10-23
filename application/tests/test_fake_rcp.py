import json

import pymupdf
import pytest
from docx import Document as DocxDocument

from scripts.fake_rcp import generate_dataset
from scripts.fake_rcp.centers import CENTERS
from scripts.fake_rcp.generator import CaseGenerator
from scripts.fake_rcp.model import FICTITIOUS_BANNER
from scripts.fake_rcp.pathologies import PATHOLOGIES


def _pdf_text(path):
    flags = pymupdf.TEXTFLAGS_TEXT & ~pymupdf.TEXT_PRESERVE_LIGATURES
    return " ".join(
        str(page.get_text(flags=flags)) for page in pymupdf.open(path)
    ).replace("\n", " ")


def _docx_text(path):
    doc = DocxDocument(path)
    cells = [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]
    return " ".join([p.text for p in doc.paragraphs] + cells)


@pytest.fixture(scope="module")
def dataset(tmp_path_factory):
    out = tmp_path_factory.mktemp("fake_rcp")
    cases = generate_dataset(out, count=len(CENTERS) * 2, seed=42)
    return out, cases


def test_one_docx_and_one_pdf_per_case_plus_manifest(dataset):
    out, cases = dataset
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert len(manifest) == len(cases)
    for entry in manifest:
        assert sorted(entry["files"]) == [
            f"{entry['case_id']}.sample.docx",
            f"{entry['case_id']}.sample.pdf",
        ]
        for name in entry["files"]:
            assert (out / name).stat().st_size > 0


def test_all_centers_are_used(dataset):
    _, cases = dataset
    assert {c.center_id for c in cases} == set(CENTERS)
    assert len({c.pathology_id for c in cases}) > 1


@pytest.mark.parametrize("reader,ext", [(_pdf_text, "pdf"), (_docx_text, "docx")])
def test_documents_are_marked_fictitious_and_contain_case_data(dataset, reader, ext):
    out, cases = dataset
    for case in cases:
        text = reader(out / f"{case.case_id}.sample.{ext}")
        assert FICTITIOUS_BANNER.split(" — ")[0] in text
        assert case.patient.last_name in text
        assert case.question in text
        for item in case.items:
            assert item.result.split(",")[0] in text


def test_post_meeting_documents_list_decision_and_participants(tmp_path):
    cases = generate_dataset(
        tmp_path, count=len(CENTERS), seed=3, stage="post", formats=("pdf",)
    )
    for case in cases:
        text = _pdf_text(tmp_path / f"{case.case_id}.sample.pdf")
        assert case.decision in text
        assert all(p.last_name in text for _, p in case.participants)


def test_generation_is_reproducible(tmp_path):
    first = generate_dataset(tmp_path / "a", count=3, seed=7, formats=("docx",))
    second = generate_dataset(tmp_path / "b", count=3, seed=7, formats=("docx",))
    assert [c.ground_truth() for c in first] == [c.ground_truth() for c in second]


def test_ground_truth_is_consistent():
    generator = CaseGenerator(seed=1, missing_rate=0.3)
    for i in range(50):
        pathology = list(PATHOLOGIES)[i % len(PATHOLOGIES)]
        case = generator.generate(f"c{i}", "chu_valcourt", pathology)
        truth = case.ground_truth()
        expected_ids = {spec.id for spec in PATHOLOGIES[pathology].items}
        assert (
            set(truth["present_items"]) | {m["id"] for m in truth["missing_items"]}
            == expected_ids
        )
        missing_required = [m for m in truth["missing_items"] if m["required"]]
        assert truth["complete"] == (
            not missing_required and not truth["outdated_items"]
        )
        kinds = [r["kind"] for r in truth["expected_requests"]]
        assert kinds.count("complement_bilan") == len(missing_required)
        assert ("oncogeriatrie" in kinds) == (case.age >= 75)


def test_missing_rate_zero_gives_every_item():
    case = CaseGenerator(seed=5, missing_rate=0).generate(
        "c", "chu_valcourt", "pancreas"
    )
    assert case.missing == []
