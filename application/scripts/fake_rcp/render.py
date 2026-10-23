"""Rendu des documents fictifs en Word (python-docx) et en PDF (PyMuPDF Story)."""

from html import escape
from pathlib import Path
from typing import cast

import pymupdf
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.styles.style import ParagraphStyle

from .model import Banner, Checklist, Document, Fields, Heading, Paragraph, Table, Title

DOCX_FONTS = {"sans": "Arial", "serif": "Times New Roman", "mono": "Courier New"}
CSS_FONTS = {"sans": "sans-serif", "serif": "serif", "mono": "monospace"}
CHECKED, UNCHECKED = "☒", "☐"


def _pairs_to_rows(pairs: list[tuple[str, str]], columns: int) -> list[list[str]]:
    rows = []
    for i in range(0, len(pairs), columns):
        row: list[str] = []
        for label, value in pairs[i : i + columns]:
            row += [label, value]
        rows.append(row + [""] * (2 * columns - len(row)))
    return rows


def _heading_text(doc: Document, text: str) -> str:
    return text.upper() if doc.style.uppercase_headings else text


# --- Word --------------------------------------------------------------------------


def _shade(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _docx_table(
    word, rows: list[list[str]], header: list[str] | None, accent: str, label_cols: bool
):
    ncols = len(header or rows[0])
    table = word.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for cell, text in zip(cells, header):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
            _shade(cell, accent)
    for row in rows:
        cells = table.add_row().cells
        for col, (cell, text) in enumerate(zip(cells, row)):
            cell.text = text
            if label_cols and col % 2 == 0 and text:
                cell.paragraphs[0].runs[0].bold = True
                _shade(cell, "F2F2F2")
    word.add_paragraph()


def to_docx(doc: Document, path: Path) -> None:
    word = DocxDocument()
    accent = RGBColor.from_string(doc.style.accent)
    normal = cast(ParagraphStyle, word.styles["Normal"])
    normal.font.name = DOCX_FONTS[doc.style.font]
    normal.font.size = Pt(10)

    section = word.sections[0]
    section.header.paragraphs[0].text = doc.header
    section.footer.paragraphs[0].text = doc.footer
    word.core_properties.title = f"Fiche RCP fictive {doc.case.case_id}"
    word.core_properties.author = "Générateur de fiches RCP fictives Oncoflow"
    word.core_properties.comments = "Document fictif : aucune donnée réelle."

    for block in doc.blocks:
        match block:
            case Banner(text=text):
                para = word.add_paragraph()
                run = para.add_run(text)
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string("C0392B")
            case Title(text=text, subtitle=subtitle):
                para = word.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(text)
                run.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = accent
                if subtitle:
                    sub = word.add_paragraph(subtitle)
                    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            case Heading(text=text):
                run = word.add_paragraph().add_run(_heading_text(doc, text))
                run.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = accent
            case Paragraph(text=text, bold=bold, italic=italic):
                run = word.add_paragraph().add_run(text)
                run.bold, run.italic = bold, italic
            case Fields(pairs=pairs, columns=columns):
                _docx_table(
                    word,
                    _pairs_to_rows(pairs, columns),
                    None,
                    doc.style.accent,
                    label_cols=True,
                )
            case Table(header=header, rows=rows):
                if rows:
                    _docx_table(word, rows, header, doc.style.accent, label_cols=False)
                else:
                    word.add_paragraph("Aucun élément transmis.")
            case Checklist(entries=entries):
                for label, checked in entries:
                    word.add_paragraph(f"{CHECKED if checked else UNCHECKED} {label}")
    word.save(str(path))


# --- PDF ---------------------------------------------------------------------------


def _html_table(
    rows: list[list[str]], header: list[str] | None, accent: str, label_cols: bool
) -> str:
    out = ["<table>"]
    if header:
        out.append(
            "<tr>"
            + "".join(
                f'<th style="background-color:#{accent}">{escape(h)}</th>'
                for h in header
            )
            + "</tr>"
        )
    for row in rows:
        cells = []
        for col, text in enumerate(row):
            cls = ' class="label"' if label_cols and col % 2 == 0 and text else ""
            cells.append(f"<td{cls}>{escape(text)}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table>")
    return "".join(out)


def to_html(doc: Document) -> str:
    accent = doc.style.accent
    parts = []
    for block in doc.blocks:
        match block:
            case Banner(text=text):
                parts.append(f'<p class="banner">{escape(text)}</p>')
            case Title(text=text, subtitle=subtitle):
                parts.append(f"<h1>{escape(text)}</h1>")
                if subtitle:
                    parts.append(f'<p class="subtitle">{escape(subtitle)}</p>')
            case Heading(text=text):
                parts.append(f"<h2>{escape(_heading_text(doc, text))}</h2>")
            case Paragraph(text=text, bold=bold, italic=italic):
                html = escape(text).replace("\n", "<br/>")
                if bold:
                    html = f"<b>{html}</b>"
                if italic:
                    html = f"<i>{html}</i>"
                parts.append(f"<p>{html}</p>")
            case Fields(pairs=pairs, columns=columns):
                parts.append(
                    _html_table(
                        _pairs_to_rows(pairs, columns), None, accent, label_cols=True
                    )
                )
            case Table(header=header, rows=rows):
                parts.append(
                    _html_table(rows, header, accent, False)
                    if rows
                    else "<p>Aucun élément transmis.</p>"
                )
            case Checklist(entries=entries):
                parts += [
                    f"<p>{CHECKED if checked else UNCHECKED} {escape(label)}</p>"
                    for label, checked in entries
                ]
    return "".join(parts)


def _css(doc: Document) -> str:
    accent = doc.style.accent
    return f"""
        * {{ font-family: {CSS_FONTS[doc.style.font]}; font-size: 10px; }}
        h1 {{ font-size: 16px; color: #{accent}; text-align: center; margin: 6px 0; }}
        h2 {{ font-size: 12px; color: #{accent}; margin: 10px 0 4px 0; }}
        p {{ margin: 3px 0; }}
        .banner {{ color: #C0392B; font-size: 8px; font-weight: bold; }}
        .subtitle {{ text-align: center; }}
        table {{ border-collapse: collapse; width: 100%; margin: 4px 0; }}
        th {{ color: white; text-align: left; padding: 3px; }}
        td {{ border: 1px solid #999999; padding: 3px; vertical-align: top; }}
        td.label {{ font-weight: bold; background-color: #F2F2F2; }}
    """


def to_pdf(doc: Document, path: Path) -> None:
    page = pymupdf.paper_rect("a4")
    body = page + (50, 60, -50, -50)
    story = pymupdf.Story(html=to_html(doc), user_css=_css(doc))
    writer = pymupdf.DocumentWriter(str(path))
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(body)
        story.draw(device)
        writer.end_page()
    writer.close()

    # En-tête, pied de page et métadonnées ajoutés après la mise en page
    pdf = pymupdf.open(str(path))
    for number, pdf_page in enumerate(pdf, start=1):
        pdf_page.insert_text((50, 35), doc.header, fontsize=8, color=(0.3, 0.3, 0.3))
        pdf_page.insert_text(
            (50, page.height - 25),
            f"{doc.footer} | page {number}/{pdf.page_count}",
            fontsize=7,
        )
    pdf.set_metadata(
        {
            "title": f"Fiche RCP fictive {doc.case.case_id}",
            "author": "Générateur de fiches RCP fictives Oncoflow",
            "subject": "Document fictif : aucune donnée réelle.",
            "creator": "oncoflow scripts.fake_rcp",
        }
    )
    pdf.saveIncr()
