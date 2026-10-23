import re
from datetime import date
from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from pseudonymizer.convert import convert
from pseudonymizer.detect import Detector, DetectorConfig
from pseudonymizer.entities import Span, resolve
from pseudonymizer.lexicon import Lexicon, preceding_cue
from pseudonymizer.pseudonymize import _shift_date, pseudonymize
from pseudonymizer.rules import cue_spans, normalize_cue, pattern_spans
from pseudonymizer.server import create_app
from pseudonymizer.workspace import Workspace

# Document entièrement fictif.
FORM = """\
FICHE RCP — CHU de Valcourt

| Nom de naissance | VAILLANT | Prénom | Maurice |
|---|---|---|---|
| Sexe | Masculin | Date de naissance | 11/10/1953 (72 ans) |
| Adresse | 48, boulevard Ollivier, 80728 Gauthier | IPP | 298277534 |

| Médecin traitant | Dr Catherine GUIBERT |
|---|---|

PATIENT: NORMAND, Richard    SEXE: M    DDN: 1955-06-30    IPP: 823061459
MEDECIN_TRAITANT: Dr Émile GERMAIN

Cher confrère, Mme Élise LEMONNIER, née le 4 février 1956, demeurant 6, chemin Riou à Pichon-la-Forêt,
est présentée en RCP. VAILLANT a reçu une chimiothérapie. Score de Child-Pugh A5.
"""


def labels_of(text, spans):
    return {(text[s.start : s.end], s.label) for s in spans}


@pytest.fixture(scope="module")
def detector():
    return Detector()


# --- Règles -------------------------------------------------------------------------


def test_normalize_cue():
    assert normalize_cue("Né(e) le :") == "ne le"
    assert normalize_cue("MEDECIN_TRAITANT") == "medecin traitant"
    assert normalize_cue("**Nom de<br>naissance**") == "nom de naissance"


def test_cues_in_tables_and_fields():
    found = labels_of(FORM, cue_spans(FORM, Lexicon().all_cues()))
    assert ("VAILLANT", "NOM") in found
    assert ("Maurice", "PRENOM") in found
    assert ("11/10/1953", "DATE_NAISSANCE") in found
    assert ("298277534", "IPP") in found
    assert ("GUIBERT", "NOM") in found and ("Catherine", "PRENOM") in found
    assert ("NORMAND", "NOM") in found and ("Richard", "PRENOM") in found
    assert ("1955-06-30", "DATE_NAISSANCE") in found
    assert ("823061459", "IPP") in found
    assert ("4 février 1956", "DATE_NAISSANCE") in found
    assert ("Pichon-la-Forêt", "VILLE") in found
    # Un libellé d'export n'est pas un nom.
    assert not any(text in {"SEXE", "DDN"} for text, _ in found)


def test_titled_person_with_accented_initials():
    found = labels_of(FORM, pattern_spans(FORM))
    assert ("Élise", "PRENOM") in found and ("LEMONNIER", "NOM") in found
    assert ("Émile", "PRENOM") in found and ("GERMAIN", "NOM") in found
    assert ("80728", "ZIP") in found and ("Gauthier", "VILLE") in found
    assert ("CHU de Valcourt", "HOPITAL") in found


def test_resolve_prefers_higher_tier_then_longer():
    text = "né le 11/10/1953"
    spans = [Span(6, 16, "DATE", "eds-pseudo"), Span(6, 16, "DATE_NAISSANCE", "repere")]
    assert [s.label for s in resolve(spans, text)] == ["DATE_NAISSANCE"]


# --- Détection complète --------------------------------------------------------------


def test_detector_masks_every_identifier(detector):
    output, _ = pseudonymize(FORM, detector.detect(FORM))
    for value in [
        "VAILLANT",
        "Maurice",
        "11/10/1953",
        "298277534",
        "GUIBERT",
        "NORMAND",
        "Richard",
        "823061459",
        "GERMAIN",
        "LEMONNIER",
        "Élise",
        "Riou",
        "Pichon-la-Forêt",
    ]:
        assert value not in output, value
    assert "Child-Pugh" in output and "chimiothérapie" in output
    # Propagation : la mention isolée « VAILLANT a reçu » est aussi masquée.
    assert "[NOM] a reçu" in output


def test_eds_only_baseline_misses_table_fields():
    spans = Detector(DetectorConfig.eds_only()).detect(FORM)
    assert "VAILLANT" not in {FORM[s.start : s.end] for s in spans}


def test_lexicon_learning(detector):
    text = "| Médecin coordonnateur | ZORGLUB |\n|---|---|\nAvis de ZORGLUB. Score de Gleason 7."
    start = text.index("ZORGLUB")
    lexicon = Lexicon()
    lexicon.learn_added(text, Span(start, start + 7, "NOM", "manuel"))
    assert lexicon.always == {"zorglub": "NOM"}
    assert lexicon.cues == {"medecin coordonnateur": "PERSONNE"}
    gleason = text.index("Gleason")
    lexicon.learn_rejected(text, Span(gleason, gleason + 7, "NOM", "eds-pseudo"))
    learned = Detector(lexicon=lexicon)
    found = labels_of(text, learned.detect(text))
    assert ("ZORGLUB", "NOM") in found
    assert not any(t == "Gleason" for t, _ in found)
    # Le libellé appris s'applique à un autre nom.
    other = "| Médecin coordonnateur | Dr Ada LOVELACE |"
    assert ("LOVELACE", "NOM") in labels_of(other, learned.detect(other))


def test_preceding_cue():
    text = "Interne : Jean DUPONT"
    assert preceding_cue(text, text.index("Jean")) == "interne"
    assert preceding_cue("Bonjour DUPONT", 8) is None


# --- Remplacement ---------------------------------------------------------------------


def test_shift_date_keeps_format():
    assert _shift_date("11/10/1953", 10) == "21/10/1953"
    assert _shift_date("1955-06-30", 1) == "1955-07-01"
    assert _shift_date("4 février 1956", 30) == "5 mars 1956"
    assert _shift_date("31/02/2026", 1) is None


def test_surrogates_are_consistent(detector):
    text = "Dr Ada LOVELACE a vu le patient. LOVELACE conclut. Revu le 01/03/2026 et le 11/03/2026."
    output, mapping = pseudonymize(text, detector.detect(text), "surrogate", seed="doc")
    replacements = {m["original"]: m["replacement"] for m in mapping}
    assert "LOVELACE" not in output
    assert output.count(replacements["LOVELACE"]) == 2
    dates = re.findall(r"\d{2}/\d{2}/\d{4}", output)
    assert len(dates) == 2 and "01/03/2026" not in dates
    first, second = (date(int(y), int(m), int(d)) for d, m, y in (x.split("/") for x in dates))
    assert (second - first).days == 10  # l'intervalle clinique est préservé


# --- Conversion, espace de travail et API --------------------------------------------


@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    path = tmp_path / "fiche.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 80), "DOCUMENT FICTIF", fontsize=12)
    page.insert_text((50, 110), "Patient : Mme Ada LOVELACE", fontsize=11)
    page.insert_text((50, 130), "Date de naissance : 10/12/1950", fontsize=11)
    page.insert_text((50, 150), "Adénocarcinome du pancréas, CA 19-9 à 37 UI/mL.", fontsize=11)
    doc.save(str(path))
    return path


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    path = tmp_path / "fiche.docx"
    word = DocxDocument()
    word.add_paragraph("DOCUMENT FICTIF")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Nom de naissance"
    table.rows[0].cells[1].text = "BABBAGE"
    word.add_paragraph("Le Dr Charles BABBAGE présente le dossier.")
    word.save(str(path))
    return path


def test_convert_pdf_with_pymupdf(sample_pdf):
    result = convert(sample_pdf, "pymupdf")
    assert result.engine == "pymupdf"
    assert "Ada LOVELACE" in result.markdown


def test_convert_docx_with_docling(sample_docx):
    result = convert(sample_docx)
    assert result.engine == "docling"
    assert "BABBAGE" in result.markdown


def test_workspace_and_api(tmp_path, sample_pdf):
    workspace = Workspace(tmp_path / "ws")
    client = TestClient(create_app(workspace))
    assert "Pseudonymiseur" in client.get("/").text

    with sample_pdf.open("rb") as f:
        doc_id = client.post("/api/documents", files={"file": ("fiche.pdf", f, "application/pdf")}).json()["id"]
    doc = client.get(f"/api/documents/{doc_id}").json()
    text = doc["markdown"]
    found = {(e["text"], e["label"]) for e in doc["entities"]}
    assert ("LOVELACE", "NOM") in found and ("10/12/1950", "DATE_NAISSANCE") in found
    assert len(doc["pages"]) == 1

    # Localisation dans le PDF d'origine
    locations = client.get(f"/api/documents/{doc_id}/locations").json()
    assert any(rects for rects in locations.values())
    assert client.get(f"/api/documents/{doc_id}/pages/0.png").headers["content-type"] == "image/png"

    # Ajout d'un identifiant oublié (« pancréas » fait office de test) et rejet d'une détection
    start = text.index("pancréas")
    lovelace = next(e for e in doc["entities"] if e["text"] == "LOVELACE")
    body = {
        "added": [{"start": start, "end": start + len("pancréas"), "label": "AUTRE"}],
        "rejected": [lovelace],
        "relabeled": [],
    }
    entities = client.put(f"/api/documents/{doc_id}/annotations", json=body).json()["entities"]
    states = {e["text"]: e["state"] for e in entities}
    assert states["pancréas"] == "added" and states["LOVELACE"] == "rejected"
    output = client.get(f"/api/documents/{doc_id}/output").text
    assert "[AUTRE]" in output and "LOVELACE" in output

    lexicon = client.get("/api/lexicon").json()
    assert lexicon["always"] == {"pancréas": "AUTRE"} and "lovelace" in lexicon["never"]

    # Annuler les corrections les retire aussi du lexique
    client.put(f"/api/documents/{doc_id}/annotations?learn=false", json={"added": [], "rejected": [], "relabeled": []})
    assert client.get("/api/lexicon").json()["always"] == {}
    client.put(f"/api/documents/{doc_id}/annotations", json=body)

    # Export au format d'entraînement d'eds-pseudo, une fois le document validé
    assert client.get("/api/export/training.jsonl").text.strip() == ""
    client.put(f"/api/documents/{doc_id}/status/valide")
    line = client.get("/api/export/training.jsonl").text.strip()
    assert '"note_text"' in line and '"label": "AUTRE"' in line


def test_upload_rejects_non_pdf(tmp_path):
    client = TestClient(create_app(Workspace(tmp_path / "ws")))
    response = client.post("/api/documents", files={"file": ("x.pdf", b"not a pdf", "application/pdf")})
    assert response.status_code == 415
    response = client.post("/api/documents", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert response.status_code == 415
    assert client.get("/api/documents/../../etc").status_code == 404
